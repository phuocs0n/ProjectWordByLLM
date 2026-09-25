"""Renderer tất định: ReportSpec + style profile -> tài liệu .docx.

Bố cục 3 section giống báo cáo mẫu:
  1. Trang bìa  - khung viền, không header/footer.
  2. Phần đầu   - LỜI MỞ ĐẦU, MỤC LỤC; footer có số trang.
  3. Nội dung   - header (Khoa | Mã lớp) có đường kẻ, footer số trang có đường kẻ.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ..numbering import CaptionCounter, HeadingNumberer
from ..spec import (
    CodeBlock,
    DividerBlock,
    FigurePlaceholderBlock,
    HeadingBlock,
    ImageBlock,
    ListBlock,
    NoteBlock,
    PageBreakBlock,
    ParagraphBlock,
    ReportSpec,
    TableBlock,
)
from ..style_profile import load_profile
from ..watermark_remover import remove_watermarks
from . import ooxml

ALIGN = {
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}

PAGE_SIZES_CM = {"letter": (21.59, 27.94), "a4": (21.0, 29.7)}

FRONT_STYLE = "Front Heading"
CODE_STYLE = "Code Block"
CAPTION_STYLE = "Caption"

# **đậm**, *nghiêng*, `mã`, [chữ](url)
# **đậm**, *nghiêng*, `mã`, [chữ](url), ~chỉ số dưới~, ^chỉ số trên^
_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\)|~[^~\s][^~]*?~|\^[^^\s][^^]*?\^)")

NOTE_LABELS = {"note": "Lưu ý: ", "tip": "Mẹo: ", "warning": "Cảnh báo: "}
NOTE_COLORS = {"note": ("2F5496", "EAF1FB"), "tip": ("2E7D32", "EDF7EE"), "warning": ("C00000", "FDECEA")}


class DocxRenderer:
    def __init__(self, profile: dict[str, Any] | str = "hcmus-clc", base_dir: str | Path = ".") -> None:
        self.p = load_profile(profile) if isinstance(profile, str) else profile
        self.base_dir = Path(base_dir)
        self.font = self.p["fonts"]["body"]
        self.body_size = self.p["body"]["size_pt"]

    # ------------------------------------------------------------------ public

    def render(self, spec: ReportSpec):
        doc = Document()
        self.doc = doc
        self._pending_break = False
        self.scheme = self.p.get("heading_numbering", "section")
        # caption_numbering: chapter ("Hình 2.3") hoặc global ("Hình 7"); mặc định theo heading_numbering
        self.caption_scheme = "chapter" if self.p.get("caption_numbering", self.scheme) == "chapter" else "section"
        self.labels = dict(self.p["labels"])
        if spec.preface_title:
            self.labels["preface"] = spec.preface_title
        if spec.references_title:
            self.labels["references"] = spec.references_title
        self.references_numbered = (self.p.get("references_numbered", False)
                                    if spec.references_numbered is None else spec.references_numbered)
        self.counter = CaptionCounter(self.caption_scheme)
        self._setup_styles()
        self.numbering = ooxml.NumberingManager(doc, self.font)

        cover = doc.sections[0]
        self._setup_page(cover)
        self._page_border(cover, cover=True)
        self._cover(spec)

        front = doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_page(front)
        self._page_border(front, cover=False)
        self._header_footer(front, spec, with_header=False)
        if self.p["page"].get("number_from_front"):
            ooxml.page_number_start(front, 1)  # bìa không đánh số, trang sau bìa là trang 1
        toc_entries = self._toc_entries(spec)
        captions = self._caption_entries(spec)
        started = False

        def new_page() -> None:
            nonlocal started
            if started:
                self._break_before_next()
            started = True

        if spec.preface:
            new_page()
            self._front_heading(self.labels["preface"])
            for text in spec.preface:
                self._paragraph(text)
            self._divider()
        if spec.front_matter:
            first = spec.front_matter[0]
            if not (isinstance(first, HeadingBlock) and first.level == 1):
                new_page()
                self._flush_break()
            started = self._blocks(spec.front_matter, self._numberer(), page_started=started)
        if spec.include_toc:
            new_page()
            toc_title = self._heading_para(self.labels["toc"], FRONT_STYLE)
            title_cfg = self.p.get("toc", {}).get("title")
            if title_cfg:
                toc_title.alignment = ALIGN[title_cfg.get("align", "center")]
                for run in toc_title.runs:
                    ooxml.set_run_font(run, self.font, title_cfg.get("size_pt", 20))
                    run.font.color.rgb = ooxml.hex_color(title_cfg.get("color", "000000"))
            self._toc(toc_entries, ' TOC \\o "1-3" \\h \\z \\u ', "toc")
            self._divider()
        for flag, kind, key in ((spec.list_of_figures, "figure", "lof"), (spec.list_of_tables, "table", "lot")):
            if flag:
                new_page()
                self._front_heading(self.labels[key])
                label = self.p["caption"][f"{kind}_label"]
                self._toc([(0, text) for text in captions[kind]], f' TOC \\h \\z \\c "{label}" ', "table of figures")

        body = doc.add_section(WD_SECTION.NEW_PAGE)
        ooxml.clear_page_number_start(body)
        self._setup_page(body)
        self._page_border(body, cover=False)
        self._header_footer(body, spec, with_header=True)
        self._body(spec)

        ooxml.enable_update_fields_on_open(doc)
        props = doc.core_properties
        props.title = " - ".join(" ".join(x.split()) for x in (spec.meta.report_type, spec.meta.subject) if x)
        props.subject = " ".join(spec.meta.topic.split())  # tên đề tài có thể xuống dòng trên bìa
        props.author = self.author(spec)
        props.comments = ""
        return doc

    @staticmethod
    def author(spec: ReportSpec) -> str:
        return ", ".join(m.name for m in spec.meta.members)

    def save(self, spec: ReportSpec, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.render(spec).save(str(path))
        remove_watermarks(path, author=self.author(spec))  # không để lại nhãn trình tạo/AI
        return path

    # ------------------------------------------------------------------ setup

    @property
    def text_width_cm(self) -> float:
        page = self.p["page"]
        width, _ = PAGE_SIZES_CM[page["size"]]
        return width - page["margin_left_cm"] - page["margin_right_cm"]

    def _page_border(self, section, cover: bool) -> None:
        ooxml.clear_page_border(section)
        page = self.p["page"]
        border = page.get("border")
        if border and (cover or border.get("all_pages")):
            ooxml.page_border(section, border.get("color", "000000"), border.get("size", 12),
                              border.get("style", "single"), border.get("space", 24))
        elif cover and page.get("cover_border"):
            ooxml.page_border(section)

    def _setup_page(self, section) -> None:
        page = self.p["page"]
        width, height = PAGE_SIZES_CM[page["size"]]
        section.page_width = Cm(width)
        section.page_height = Cm(height)
        section.top_margin = Cm(page["margin_top_cm"])
        section.bottom_margin = Cm(page["margin_bottom_cm"])
        section.left_margin = Cm(page["margin_left_cm"])
        section.right_margin = Cm(page["margin_right_cm"])
        section.header_distance = Cm(page["header_distance_cm"])
        section.footer_distance = Cm(page["footer_distance_cm"])

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        body = self.p["body"]
        normal = styles["Normal"]
        ooxml.set_style_font(normal, self.font, self.body_size)
        pf = normal.paragraph_format
        pf.line_spacing = body["line_spacing"]
        pf.space_after = Pt(body["space_after_pt"])
        pf.space_before = Pt(0)

        head = self.p["headings"]
        for level in (1, 2, 3):
            cfg = {"color": head["color"], **head[f"level{level}"]}
            style = styles[f"Heading {level}"]
            ooxml.set_style_font(style, self.font, cfg["size_pt"])
            style.font.bold = cfg.get("bold", True)
            style.font.italic = cfg.get("italic", False)
            style.font.color.rgb = ooxml.hex_color(cfg["color"])
            spf = style.paragraph_format
            spf.space_before = Pt(cfg["space_before_pt"])
            spf.space_after = Pt(cfg["space_after_pt"])
            spf.left_indent = Cm(cfg.get("indent_cm", 0))
            spf.first_line_indent = Cm(0)
            spf.keep_with_next = True
            spf.line_spacing = 1.0
            spf.alignment = WD_ALIGN_PARAGRAPH.LEFT

        front_cfg = head["front"]
        front = styles.add_style(FRONT_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        front.base_style = styles["Heading 1"]
        front.next_paragraph_style = normal
        ooxml.set_style_font(front, self.font, front_cfg["size_pt"])
        front.font.bold = True
        front.font.color.rgb = ooxml.hex_color(front_cfg["color"])
        front.paragraph_format.alignment = ALIGN[front_cfg["align"]]
        front.paragraph_format.space_before = Pt(0)
        front.paragraph_format.space_after = Pt(18)
        front.paragraph_format.left_indent = Cm(0)
        ooxml.set_outline_level(front, 0)  # để MỤC LỤC liệt kê được

        cap = self.p["caption"]
        caption = styles[CAPTION_STYLE] if CAPTION_STYLE in [s.name for s in styles] else styles.add_style(CAPTION_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        caption.base_style = normal
        ooxml.set_style_font(caption, self.font, cap["size_pt"])
        caption.font.italic = cap["italic"]
        caption.font.bold = False
        caption.font.color.rgb = ooxml.hex_color("000000")
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(10)

        code_cfg = self.p["code"]
        code = styles.add_style(CODE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        ooxml.set_style_font(code, self.p["fonts"]["code"], code_cfg["size_pt"])
        code.paragraph_format.line_spacing = 1.0
        code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        code.paragraph_format.space_after = Pt(8)
        code.paragraph_format.keep_together = True

        tof = styles.add_style("table of figures", WD_STYLE_TYPE.PARAGRAPH)
        tof.base_style = normal
        tof.paragraph_format.space_after = Pt(3)
        tof.paragraph_format.first_line_indent = Cm(0)
        tof.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tof.paragraph_format.tab_stops.add_tab_stop(Cm(self.text_width_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

        for level in (1, 2, 3):
            toc = styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
            toc.base_style = normal
            toc.paragraph_format.left_indent = Cm(0.6 * (level - 1))
            toc.paragraph_format.space_after = Pt(3)
            toc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            toc.paragraph_format.tab_stops.add_tab_stop(
                Cm(self.text_width_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
            )
            toc.font.bold = level == 1 and self.p.get("toc", {}).get("bold_level1", True)
            toc.paragraph_format.line_spacing = self.p.get("toc", {}).get("line_spacing", toc.paragraph_format.line_spacing)

    # ------------------------------------------------------------------ primitives

    def _inline(self, paragraph, text: str, size_pt: float | None = None, bold: bool | None = None,
                italic: bool | None = None) -> None:
        size_pt = size_pt or self.body_size
        for part in _INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                run = paragraph.add_run(part[2:-2])
                run.bold = True
                ooxml.set_run_font(run, self.font, size_pt)
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                run = paragraph.add_run(part[1:-1])
                ooxml.set_run_font(run, self.p["fonts"]["code"], size_pt - 1)
            elif part.startswith("[") and "](" in part and part.endswith(")"):
                label, url = part[1:-1].split("](", 1)
                ooxml.add_hyperlink(paragraph, label, url, self.font, size_pt, self.p["hyperlink_color"])
                continue
            elif (part[0], part[-1]) in {("~", "~"), ("^", "^")} and len(part) > 2:
                run = paragraph.add_run(part[1:-1])
                ooxml.set_run_font(run, self.font, size_pt)
                if part[0] == "~":
                    run.font.subscript = True
                else:
                    run.font.superscript = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                run = paragraph.add_run(part[1:-1])
                run.italic = True
                ooxml.set_run_font(run, self.font, size_pt)
            else:
                run = paragraph.add_run(part)
                ooxml.set_run_font(run, self.font, size_pt)
            if bold is not None:
                run.bold = bold or run.bold
            if italic is not None:
                run.italic = italic or run.italic

    def _paragraph(self, text: str, align: str | None = None, container=None):
        container = container or self.doc
        para = container.add_paragraph()
        para.alignment = ALIGN[align or self.p["body"]["align"]]
        indent = self.p["body"].get("first_line_indent_cm", 0)
        if indent:
            para.paragraph_format.first_line_indent = Cm(indent)
        self._inline(para, text)
        return para

    def _centered(self, text: str, size_pt: float, bold: bool = True, space_after: float = 0,
                  space_before: float = 0, italic: bool = False):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(space_before)
        para.paragraph_format.line_spacing = 1.0
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        ooxml.set_run_font(run, self.font, size_pt)
        return para

    def _front_heading(self, text: str) -> None:
        self._heading_para(text, FRONT_STYLE)

    def _heading_para(self, text: str, style: str):
        """Tiêu đề; nếu có yêu cầu sang trang thì dùng 'page break before' thay cho đoạn ngắt trang riêng,
        tránh trang trắng khi trang trước vừa đầy."""
        para = self.doc.add_paragraph(text, style=style)
        if self._pending_break:
            para.paragraph_format.page_break_before = True
            self._pending_break = False
        return para

    def _break_before_next(self) -> None:
        self._pending_break = True

    def _flush_break(self) -> None:
        if self._pending_break:
            self._pending_break = False
            self._page_break()

    def _divider(self) -> None:
        if not self.p.get("divider_text"):
            return
        self._centered(self.p["divider_text"], 12, bold=False, space_before=6, space_after=6)

    def _page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _caption_label(self, kind: str) -> str:
        return self.p["caption"][f"{kind}_label"]

    def _caption_cfg(self, kind: str) -> dict:
        base = self.p["caption"]
        return {**base, **(base.get(kind) or {})}

    def _caption(self, kind: str, text: str) -> None:
        """Chú thích 'Hình 2.3. …' / 'Bảng 1: …' - số thứ tự là field SEQ để Word lập danh mục hình/bảng.
        Định dạng theo từng loại (caption.table / caption.figure trong profile)."""
        cfg = self._caption_cfg(kind)
        label = self._caption_label(kind)
        prefix, number, first = self.counter.next(kind)
        para = self.doc.add_paragraph(style=CAPTION_STYLE)
        size = cfg["size_pt"]
        label_bold = cfg.get("label_bold", True)
        label_italic = cfg.get("label_italic", False)
        run = para.add_run(f"{label} {prefix}")
        run.bold, run.italic = label_bold, label_italic
        ooxml.set_run_font(run, self.font, size)
        reset = " \\r 1" if prefix and first else ""
        ooxml.add_field(para, f"SEQ {label}{reset} \\* ARABIC", number, font=self.font, size_pt=size, bold=label_bold)
        if label_italic:
            for r in para.runs:
                r.italic = True
        sep = para.add_run(cfg.get("separator", ": ") if text else "")
        sep.bold, sep.italic = label_bold, label_italic
        ooxml.set_run_font(sep, self.font, size)
        if text:
            self._inline(para, text, size_pt=size, bold=True if cfg.get("bold_text") else None,
                         italic=cfg.get("italic"))
        if kind == "table" and self.p["caption"].get("table_position", "above") != "below":
            para.paragraph_format.keep_with_next = True

    def _caption_entries(self, spec: ReportSpec) -> dict[str, list[str]]:
        """Văn bản chú thích theo đúng thứ tự/số sẽ render - dùng cho DANH MỤC HÌNH/BẢNG."""
        numberer, counter = self._numberer(), CaptionCounter(self.caption_scheme)
        sep = self.p["caption"].get("separator", ": ")
        out: dict[str, list[str]] = {"figure": [], "table": []}
        for block in spec.body:
            kind = None
            if isinstance(block, HeadingBlock) and block.numbered:
                numberer.next(block.level)
                if block.level == 1:
                    counter.new_chapter(numberer.chapter)
            elif isinstance(block, (ImageBlock, FigurePlaceholderBlock)):
                kind = "figure"
            elif isinstance(block, TableBlock) and block.caption:
                kind = "table"
            if kind:
                prefix, number, _ = counter.next(kind)
                caption = re.sub(r"\*\*|\*|`", "", block.caption)
                out[kind].append(f"{self._caption_label(kind)} {prefix}{number}{sep if caption else ''}{caption}")
        return out

    # ------------------------------------------------------------------ cover

    def _cover(self, spec: ReportSpec) -> None:
        for image in spec.meta.cover_images:
            path = self._resolve(image.path)
            if not path.exists():
                continue
            anchor_para = self.doc.paragraphs[0] if self.doc.paragraphs else self.doc.add_paragraph()
            run = anchor_para.add_run()
            run.add_picture(str(path), width=Cm(image.width_cm),
                            height=Cm(image.height_cm) if image.height_cm else None)
            ooxml.float_picture(run, Cm(image.x_cm), Cm(image.y_cm))
        layout = self.p.get("cover_layout")
        if layout == "banner":
            self._cover_banner(spec)
        elif layout == "title_block":
            self._cover_title_block(spec)
        else:
            self._cover_classic(spec)

    def _cover_title_block(self, spec: ReportSpec) -> None:
        """Bìa: logo + tên trường (có dòng hoa văn), tên đề tài chữ lớn nhiều dòng (xuống dòng bằng \\n
        trong meta.topic), bảng STT / HỌ VÀ TÊN / MSSV, nơi – thời gian."""
        m = spec.meta
        cfg = self.p.get("cover", {})
        top = self.doc.paragraphs[0] if self.doc.paragraphs else self.doc.add_paragraph()
        top.paragraph_format.space_after = Pt(cfg.get("top_space_pt", 30))
        head = self.doc.add_table(rows=1, cols=2)
        head.alignment = WD_TABLE_ALIGNMENT.CENTER
        ooxml.table_no_borders(head)
        widths = (Cm(cfg.get("logo_col_cm", 6.2)), Cm(self.text_width_cm - cfg.get("logo_col_cm", 6.2)))
        for c, column in enumerate(head.columns):
            column.width = widths[c]
        left, right = head.rows[0].cells
        left.width, right.width = widths
        left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        logo = self._resolve(m.logo_path) if m.logo_path else None
        if logo and logo.exists():
            lp = left.paragraphs[0]
            lp.paragraph_format.first_line_indent = None
            lp.add_run().add_picture(str(logo), width=Cm(cfg.get("logo_width_cm", 6.1)))
        lines = [x.upper() for x in (m.university, m.school, m.faculty) if x]
        if cfg.get("divider"):
            lines.append(cfg["divider"])
        for i, line in enumerate(lines):
            para = right.paragraphs[0] if i == 0 else right.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = None
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.15
            run = para.add_run(line)
            run.bold = cfg.get("header_bold", False)
            ooxml.set_run_font(run, self.font, cfg.get("header_size_pt", 13))

        title_lines = [x for x in (m.topic or m.subject).split("\n") if x.strip()]
        for i, line in enumerate(title_lines):
            para = self._centered(line.upper(), cfg.get("title_size_pt", 36),
                                  space_before=cfg.get("title_space_pt", 100) if i == 0 else 0, space_after=0)
            para.paragraph_format.line_spacing = cfg.get("title_line_spacing", 1.5)
        if m.subtitle:
            self._centered(m.subtitle, 14, italic=True, space_before=6)

        if m.members:
            spacer = self.doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(cfg.get("members_space_pt", 80))
            spacer.paragraph_format.space_after = Pt(0)
            size = cfg.get("members_size_pt", 14)
            header = cfg.get("members_columns", ["STT", "HỌ VÀ TÊN", "MSSV"])
            table = self.doc.add_table(rows=1 + len(m.members), cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            ooxml.table_borders(table)
            ratio = cfg.get("members_widths", [1, 6, 3])
            widths = [Cm(self.text_width_cm * r / sum(ratio)) for r in ratio]
            for c, column in enumerate(table.columns):
                column.width = widths[c]
            rows = [header] + [[str(i), mem.name, mem.student_id] for i, mem in enumerate(m.members, start=1)]
            for r, values in enumerate(rows):
                for c, value in enumerate(values):
                    cell = table.rows[r].cells[c]
                    cell.width = widths[c]
                    para = cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.first_line_indent = None
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing = 1.0
                    run = para.add_run(value)
                    run.bold = r == 0
                    ooxml.set_run_font(run, self.font, size)

        place = " – ".join(x for x in (m.city, m.year) if x)
        if place:
            self._centered(place, cfg.get("place_size_pt", 13), bold=cfg.get("place_bold", True),
                           space_before=cfg.get("place_space_pt", 50))

    def _cover_banner(self, spec: ReportSpec) -> None:
        """Bìa kiểu khoa ĐT-VT: logo trái + tên trường/khoa phải, tiêu đề, bảng thông tin có viền."""
        m = spec.meta
        labels = self.p["labels"]
        top = self.doc.paragraphs[0] if self.doc.paragraphs else self.doc.add_paragraph()
        top.paragraph_format.space_after = Pt(0)
        logo = self._resolve(m.logo_path) if m.logo_path else None
        head = self.doc.add_table(rows=1, cols=2)
        head.alignment = WD_TABLE_ALIGNMENT.CENTER
        ooxml.table_no_borders(head)
        left, right = head.rows[0].cells
        widths = (Cm(5.8), Cm(self.text_width_cm - 5.8))
        for c, column in enumerate(head.columns):
            column.width = widths[c]
        left.width, right.width = widths
        left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if logo and logo.exists():
            lp = left.paragraphs[0]
            lp.paragraph_format.first_line_indent = None
            lp.add_run().add_picture(str(logo), width=Cm(5.6))
        lines = [x for x in (m.university, m.school, m.faculty) if x]
        for i, line in enumerate(lines):
            para = right.paragraphs[0] if i == 0 else right.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.first_line_indent = None
            run = para.add_run(line.upper())
            run.bold = True
            ooxml.set_run_font(run, self.font, 13)

        self._centered(m.report_type.upper(), 15, space_before=30, space_after=0)
        title = m.topic or m.subject
        self._centered(title.upper(), 17, space_before=30, space_after=6)
        if m.subtitle:
            self._centered(m.subtitle, 13, italic=True, space_before=6)

        rows: list[tuple[str, list[str]]] = []
        if m.subject and m.topic:
            rows.append((labels["subject"], [m.subject]))
        if m.instructors:
            rows.append((labels["instructors"], m.instructors))
        if m.members:
            rows.append((labels["members"], [
                f"{mem.name} – {mem.student_id}" if mem.student_id else mem.name for mem in m.members]))
        if m.class_code:
            rows.append((labels["class"], [m.class_code]))
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(60)
        if rows:
            table = self.doc.add_table(rows=len(rows), cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            ooxml.table_borders(table)
            widths = (Cm(5.2), Cm(self.text_width_cm - 5.2))
            for c, column in enumerate(table.columns):
                column.width = widths[c]
            for r, (label, values) in enumerate(rows):
                lcell, rcell = table.rows[r].cells
                lcell.width, rcell.width = widths
                lcell.vertical_alignment = rcell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                lp = lcell.paragraphs[0]
                lp.paragraph_format.space_before = lp.paragraph_format.space_after = Pt(6)
                lp.paragraph_format.first_line_indent = None
                run = lp.add_run(label.rstrip(":"))
                run.bold = True
                ooxml.set_run_font(run, self.font, 12)
                for i, value in enumerate(values):
                    vp = rcell.paragraphs[0] if i == 0 else rcell.add_paragraph()
                    vp.paragraph_format.first_line_indent = None
                    vp.paragraph_format.space_before = Pt(6 if i == 0 else 0)
                    vp.paragraph_format.space_after = Pt(6 if i == len(values) - 1 else 0)
                    vp.paragraph_format.line_spacing = 1.5
                    ooxml.set_run_font(vp.add_run(value), self.font, 12)
        place = " – ".join(x for x in (m.city, m.year) if x)
        if place:
            self._centered(place.upper(), 13, bold=False, italic=True, space_before=40)

    def _cover_classic(self, spec: ReportSpec) -> None:
        m = spec.meta
        labels = self.p["labels"]
        if m.university:
            self._centered(m.university.upper(), 14, space_after=4)
        if m.school:
            self._centered(m.school.upper(), 14, space_after=4 if self.p.get("cover_show_faculty") else 12)
        if m.faculty and self.p.get("cover_show_faculty"):
            self._centered(m.faculty.upper(), 14, space_after=12)
        logo = self._resolve(m.logo_path) if m.logo_path else None
        if logo and logo.exists():
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(12)
            para.add_run().add_picture(str(logo), width=Cm(5))
            gap = 36
        else:
            gap = 120
        self._centered(m.report_type.upper(), 24, space_before=gap, space_after=0)
        if m.subject:
            self._centered(m.subject.upper(), 24, space_after=18)
        if m.topic:
            para = self._centered("", 16, space_after=24)
            run = para.add_run(f"{labels['topic']} ")
            run.bold = True
            ooxml.set_run_font(run, self.font, 16)
            run = para.add_run(m.topic.upper())
            run.bold = True
            ooxml.set_run_font(run, self.font, 16)

        rows: list[tuple[str, list[str]]] = []
        if m.instructors:
            rows.append((labels["instructors"], m.instructors))
        if m.members:
            rows.append((labels["members"], [
                f"{mem.name} – {mem.student_id}" if mem.student_id else mem.name for mem in m.members
            ]))
        if m.class_code and self.p.get("cover_show_class"):
            rows.append((labels["class"], [m.class_code]))
        if rows:
            table = self.doc.add_table(rows=len(rows), cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            ooxml.table_no_borders(table)
            widths = (Cm(6.5), Cm(8.5))
            for r, (label, values) in enumerate(rows):
                left, right = table.rows[r].cells
                left.width, right.width = widths
                lp = left.paragraphs[0]
                lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                lp.paragraph_format.space_after = Pt(12)
                run = lp.add_run(label)
                run.bold = True
                ooxml.set_run_font(run, self.font, 14)
                for i, value in enumerate(values):
                    vp = right.paragraphs[0] if i == 0 else right.add_paragraph()
                    vp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    vp.paragraph_format.space_after = Pt(12 if i == len(values) - 1 else 0)
                    ooxml.set_run_font(vp.add_run(value), self.font, 14)

        place = " – ".join(x for x in (m.city, m.year) if x)
        if place:
            # Đẩy dòng "TP. ... – năm" xuống gần đáy trang bìa
            lines_used = len(m.instructors) + len(m.members) + (2 if m.topic else 0)
            if logo and logo.exists():
                lines_used += 8
            self._centered(place, 14, space_before=max(36, 300 - 17 * lines_used))

    def _resolve(self, path: str) -> Path:
        candidate = Path(path)
        return candidate if candidate.is_absolute() else self.base_dir / candidate

    # ------------------------------------------------------------------ header / footer

    def _header_footer(self, section, spec: ReportSpec, with_header: bool) -> None:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header_cfg, footer_cfg = self.p["header"], self.p["footer"]
        values = spec.meta.model_dump()

        hp = section.header.paragraphs[0]
        # Style "Header" mặc định có tab giữa 8.25cm làm lệch chữ bên phải -> dùng Normal
        hp.style = self.doc.styles["Normal"]
        hp.paragraph_format.space_after = Pt(0)
        left = header_cfg.get("left", "").format(**values)
        right = header_cfg.get("right", "").format(**values)
        if with_header and (left or right):
            hp.paragraph_format.tab_stops.add_tab_stop(Cm(self.text_width_cm), WD_TAB_ALIGNMENT.RIGHT)
            ooxml.set_run_font(hp.add_run(f"{left}\t{right}"), self.font, header_cfg["size_pt"])
            if header_cfg.get("border_bottom"):
                ooxml.paragraph_border(hp, {"bottom": {"sz": "6", "color": "000000"}})

        fp = section.footer.paragraphs[0]
        fp.style = self.doc.styles["Normal"]
        fp.paragraph_format.space_after = Pt(0)
        if footer_cfg.get("page_number", "center") == "none":
            return
        fp.alignment = ALIGN[footer_cfg.get("page_number", "center")]
        ooxml.add_field(fp, "PAGE", "1", font=footer_cfg.get("font", self.font), size_pt=footer_cfg["size_pt"])
        if footer_cfg.get("border_top"):
            ooxml.paragraph_border(fp, {"top": {"sz": "6", "color": "000000"}})

    # ------------------------------------------------------------------ TOC

    def _numberer(self) -> HeadingNumberer:
        return HeadingNumberer(self.scheme, self.p["headings"].get("level1_format", "{roman}."))

    def _heading(self, block: HeadingBlock, numberer: HeadingNumberer) -> tuple[str, str]:
        """(style, text) cho một tiêu đề; dùng chung cho thân bài và MỤC LỤC để luôn khớp nhau."""
        if not block.numbered:
            return (FRONT_STYLE if block.level == 1 else f"Heading {block.level}"), block.text
        return f"Heading {block.level}", f"{numberer.next(block.level)} {block.text}"

    def _toc_entries(self, spec: ReportSpec) -> list[tuple[int, str]]:
        entries: list[tuple[int, str]] = []
        if spec.preface:
            entries.append((1, self.labels["preface"]))
        front_numberer = self._numberer()
        for block in spec.front_matter:
            if isinstance(block, HeadingBlock):
                entries.append((block.level, self._heading(block, front_numberer)[1]))
        if spec.include_toc and self.p.get("toc", {}).get("include_self", True):
            entries.append((1, self.labels["toc"]))
        if spec.list_of_figures:
            entries.append((1, self.labels["lof"]))
        if spec.list_of_tables:
            entries.append((1, self.labels["lot"]))
        numberer = self._numberer()
        for block in spec.body:
            if isinstance(block, HeadingBlock):
                entries.append((block.level, self._heading(block, numberer)[1]))
        if spec.references:
            title = self.labels["references"]
            if self.references_numbered:
                title = f"{numberer.next(1)} {title}"
            entries.append((1, title))
        return entries

    def _toc(self, entries: list[tuple[int, str]], instr_text: str, style_prefix: str) -> None:
        """Field TOC (mục lục hoặc danh mục hình/bảng) có sẵn kết quả đệm không số trang;
        postprocess.update_toc() hoặc Word (F9) điền số trang thật."""

        def style(level: int) -> str:
            return style_prefix if style_prefix == "table of figures" else f"{style_prefix} {max(level, 1)}"

        first = self.doc.add_paragraph(style=style(entries[0][0] if entries else 1))

        # begin + instr + separate nằm trong đoạn đầu tiên, end nằm trong đoạn cuối
        def fld(paragraph, kind: str):
            run = paragraph.add_run()
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._element.append(el)

        fld(first, "begin")
        instr = first.add_run()
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = instr_text
        instr._element.append(instr_el)
        fld(first, "separate")
        paragraphs = []
        for i, (level, text) in enumerate(entries):
            para = first if i == 0 else self.doc.add_paragraph(style=style(level))
            ooxml.set_run_font(para.add_run(f"{text}\t"), self.font, self.p.get("toc", {}).get("size_pt", self.body_size))
            paragraphs.append(para)
        if not paragraphs:
            paragraphs.append(first)
        fld(paragraphs[-1], "end")

    # ------------------------------------------------------------------ body

    def _body(self, spec: ReportSpec) -> None:
        numberer = self._numberer()
        chapter_break = self.p.get("chapter_page_break", True)
        seen_chapter = self._blocks(spec.body, numberer, page_started=False)
        self._references(spec, numberer, chapter_break and seen_chapter)

    def _blocks(self, blocks, numberer: HeadingNumberer, page_started: bool) -> bool:
        """Render danh sách block; chương (tiêu đề cấp 1) bắt đầu trang mới. Trả về đã có nội dung hay chưa."""
        chapter_break = self.p.get("chapter_page_break", True)
        for block in blocks:
            if isinstance(block, HeadingBlock):
                if block.level == 1 and page_started and chapter_break:
                    self._break_before_next()
                style, text = self._heading(block, numberer)
                if block.level == 1 and block.numbered:
                    self.counter.new_chapter(numberer.chapter)
                self._heading_para(text, style)
            elif isinstance(block, ParagraphBlock):
                self._paragraph(block.text, block.align)
            elif isinstance(block, ListBlock):
                self._list(block)
            elif isinstance(block, TableBlock):
                self._table(block)
            elif isinstance(block, ImageBlock):
                self._image(block)
            elif isinstance(block, FigurePlaceholderBlock):
                self._placeholder(block)
            elif isinstance(block, CodeBlock):
                self._code(block)
            elif isinstance(block, NoteBlock):
                self._note(block)
            elif isinstance(block, PageBreakBlock):
                self._page_break()
            elif isinstance(block, DividerBlock):
                self._divider()
            page_started = True
        return page_started

    def _references(self, spec: ReportSpec, numberer: HeadingNumberer, page_break: bool) -> None:
        if not spec.references:
            return
        if page_break:
            self._break_before_next()
        title = self.labels["references"]
        if self.references_numbered:
            self._heading_para(f"{numberer.next(1)} {title}", "Heading 1")
        else:
            self._front_heading(title)
        num_id = self.numbering.new_list(self.p.get("references_style", "bracket"))
        for ref in spec.references:
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self.numbering.apply(para, num_id, 0)
            self._inline(para, ref.text)
            if ref.url:
                if ref.text:
                    ooxml.set_run_font(para.add_run(" "), self.font, self.body_size)
                ooxml.add_hyperlink(para, ref.url, ref.url, self.font, self.body_size, self.p["hyperlink_color"])

    def _list(self, block: ListBlock, container=None) -> None:
        container = container or self.doc
        num_id = self.numbering.new_list(block.style)
        for item in block.items:
            para = container.add_paragraph()
            para.alignment = ALIGN[self.p["body"]["align"]]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.first_line_indent = None
            self.numbering.apply(para, num_id, block.level)
            self._inline(para, item)

    def _column_widths(self, block: TableBlock) -> list[float]:
        n = len(block.columns)
        weights = list(block.col_widths[:n]) if len(block.col_widths) >= n else []
        if not weights:
            for c in range(n):
                cells = [block.columns[c]] + [row[c] if c < len(row) else "" for row in block.rows]
                longest = max(max((len(line) for line in cell.split("\n")), default=0) for cell in cells)
                weights.append(min(max(longest, 4), 40))
        total = sum(weights) or 1
        width = min(self.text_width_cm, self.p["table"].get("width_cm") or self.text_width_cm)
        return [width * w / total for w in weights]

    def _table(self, block: TableBlock) -> None:
        cfg = self.p["table"]
        size = cfg["font_size_pt"]
        caption_below = self.p["caption"].get("table_position", "above") == "below"
        if block.caption and not caption_below:
            self._caption("table", block.caption)
            self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        n = len(block.columns)
        table = self.doc.add_table(rows=1 + len(block.rows), cols=n)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        ooxml.table_borders(table, cfg["border_color"])
        widths = self._column_widths(block)

        short_col = []
        for c in range(n):
            values = [row[c] if c < len(row) else "" for row in block.rows]
            short_col.append(all(len(v) <= 25 and "\n" not in v for v in values))

        for c, title in enumerate(block.columns):
            cell = table.rows[0].cells[c]
            if cfg.get("header_fill"):
                ooxml.cell_shading(cell, cfg["header_fill"])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pad = Pt(self.p["table"].get("cell_padding_pt", 0))
            para.paragraph_format.space_before = para.paragraph_format.space_after = pad
            para.paragraph_format.first_line_indent = None
            self._inline(para, title, size_pt=size, bold=True)
        ooxml.repeat_header_row(table.rows[0])

        for r, row in enumerate(block.rows, start=1):
            ooxml.cant_split_row(table.rows[r])
            for c in range(n):
                value = row[c] if c < len(row) else ""
                cell = table.rows[r].cells[c]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell_lines = value.split("\n")
                pad = Pt(self.p["table"].get("cell_padding_pt", 0))
                for i, line in enumerate(cell_lines):
                    para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                    para.paragraph_format.space_before = pad if i == 0 else Pt(0)
                    para.paragraph_format.space_after = pad if i == len(cell_lines) - 1 else Pt(0)
                    para.paragraph_format.first_line_indent = None
                    if line.startswith("- "):
                        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        para.paragraph_format.left_indent = Cm(0.4)
                        para.paragraph_format.first_line_indent = Cm(-0.4)
                        self._inline(para, "- " + line[2:], size_pt=size)
                    else:
                        body_align = self.p["table"].get("body_align", "auto")
                        centered = body_align == "center" or (body_align == "auto" and short_col[c])
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
                        self._inline(para, line, size_pt=size)

        # Đặt cả gridCol lẫn tcW để độ rộng cột có hiệu lực ở mọi trình xem
        for c, column in enumerate(table.columns):
            column.width = Cm(widths[c])
        for row in table.rows:
            for c, cell in enumerate(row.cells):
                cell.width = Cm(widths[c])
        if block.caption and caption_below:
            for cell in table.rows[-1].cells:  # giữ bảng liền với chú thích bên dưới
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True
            self._caption("table", block.caption)
            return
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def _image(self, block: ImageBlock) -> None:
        path = self._resolve(block.path)
        if not path.exists():
            self._placeholder(FigurePlaceholderBlock(type="figure_placeholder", caption=block.caption, description=f"Không tìm thấy ảnh: {block.path}"))
            return
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.first_line_indent = None
        width = min(block.width_cm, self.text_width_cm)
        para.add_run().add_picture(str(path), width=Cm(width))
        self._caption("figure", block.caption)

    def _placeholder(self, block: FigurePlaceholderBlock) -> None:
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.first_line_indent = None
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.left_indent = Cm(2)
        para.paragraph_format.right_indent = Cm(2)
        dashed = {"val": "dashed", "sz": "8", "color": "7F7F7F", "space": "18"}
        ooxml.paragraph_border(para, {side: dashed for side in ("top", "left", "bottom", "right")})
        ooxml.paragraph_shading(para, "F7F7F7")
        run = para.add_run("[ CHÈN HÌNH ]")
        run.bold = True
        ooxml.set_run_font(run, self.font, 12)
        if block.description:
            run.add_break()
            desc = para.add_run(block.description)
            desc.italic = True
            ooxml.set_run_font(desc, self.font, 12)
        self._caption("figure", block.caption)

    def _code_caption(self, block: CodeBlock, below: bool) -> None:
        cap = self.doc.add_paragraph(style=CAPTION_STYLE)
        cfg = self._caption_cfg("code")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER if below else WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.keep_with_next = not below
        cap.paragraph_format.space_before = Pt(4 if below else 0)
        cap.paragraph_format.space_after = Pt(10 if below else 2)
        self._inline(cap, block.caption, size_pt=cfg["size_pt"], italic=cfg.get("italic"),
                     bold=True if cfg.get("bold_text") else None)

    def _code(self, block: CodeBlock) -> None:
        below = self.p["code"].get("caption_position", "above") == "below"
        if block.caption and not below:
            self._code_caption(block, below=False)
        self._code_body(block, keep_with_caption=bool(block.caption and below))
        if block.caption and below:
            self._code_caption(block, below=True)

    def _code_body(self, block: CodeBlock, keep_with_caption: bool = False) -> None:
        cfg = self.p["code"]
        para = self.doc.add_paragraph(style=CODE_STYLE)
        lines = block.code.rstrip("\n").split("\n")
        if len(lines) > 30:
            # Khối dài hơn phần lớn một trang: cho phép ngắt trang, nếu không cả khối bị đẩy sang trang
            # sau và để lại chú thích lẻ loi cuối trang trước.
            para.paragraph_format.keep_together = False
        for i, line in enumerate(lines):
            run = para.add_run(line)
            ooxml.set_run_font(run, self.p["fonts"]["code"], cfg["size_pt"])
            if i < len(lines) - 1:
                run.add_break()
        if cfg.get("border", "left") == "box":
            box = {"sz": "4", "color": cfg["border_color"], "space": "4"}
            ooxml.paragraph_border(para, {side: box for side in ("top", "left", "bottom", "right")})
        else:
            ooxml.paragraph_border(para, {"left": {"sz": "18", "color": cfg["border_color"], "space": "6"}})
        ooxml.paragraph_shading(para, cfg["fill"])
        para.paragraph_format.keep_with_next = keep_with_caption

    def _note(self, block: NoteBlock) -> None:
        border, fill = NOTE_COLORS[block.kind]
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.first_line_indent = None
        ooxml.paragraph_border(para, {"left": {"sz": "18", "color": border, "space": "6"}})
        ooxml.paragraph_shading(para, fill)
        label = para.add_run(NOTE_LABELS[block.kind])
        label.bold = True
        ooxml.set_run_font(label, self.font, self.body_size)
        self._inline(para, block.text)


def render_spec(spec: ReportSpec, output: str | Path, base_dir: str | Path = ".") -> Path:
    return DocxRenderer(spec.profile, base_dir=base_dir).save(spec, output)
