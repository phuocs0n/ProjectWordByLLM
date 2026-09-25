"""Đọc & kiểm tra (lint) tài liệu Word.

- `docx_to_markdown`: chuyển .docx thành Markdown để LLM đọc/chỉnh sửa lại.
- `read_source`: đọc tài liệu nguồn bất kỳ (docx nội bộ; pdf/pptx/xlsx/html qua MarkItDown nếu có).
- `lint_docx`: phát hiện lỗi định dạng phổ biến của báo cáo sinh viên/doanh nghiệp.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .style_profile import load_profile
from .watermark_remover import find_watermarks

# Lỗi chính tả hay gặp (nhiều lỗi lấy từ chính báo cáo mẫu).
COMMON_TYPOS = {
    "Hồ Chí Mình": "Hồ Chí Minh",
    "chúng tỏ": "chứng tỏ",
    "hoạt đồng": "hoạt động",
    "đảm bào": "đảm bảo",
    "thự hiện": "thực hiện",
    "tên mien": "tên miền",
    "sự liên quan đặc biệt đối với": "sự quan tâm đặc biệt đến",
    "cài đặc": "cài đặt",
    "sử lý": "xử lý",
    "xử dụng": "sử dụng",
    "chuẩn đoán": "chẩn đoán",
}


def iter_block_items(document) -> Iterator[Paragraph | Table]:
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    name = paragraph.style.name if paragraph.style is not None else ""
    match = re.match(r"Heading (\d)", name)
    if match:
        return int(match.group(1))
    if name in {"Title", "Front Heading"}:
        return 1
    return None


def _has_image(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.findall(".//" + qn("w:drawing")))


def _is_list_item(paragraph: Paragraph) -> bool:
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        return True
    name = paragraph.style.name if paragraph.style is not None else ""
    return name.startswith("List")


def _is_caption(paragraph: Paragraph, labels: tuple[str, ...]) -> bool:
    name = paragraph.style.name if paragraph.style is not None else ""
    return name == "Caption" or paragraph.text.strip().startswith(labels)


def docx_to_markdown(path: str | Path) -> str:
    document = Document(str(path))
    out: list[str] = []
    for item in iter_block_items(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            level = _heading_level(item)
            if _has_image(item):
                out.append("![hình ảnh]()")
            if not text:
                continue
            if level:
                out.append(f"{'#' * level} {text}")
            elif _is_list_item(item):
                out.append(f"- {text}")
            else:
                out.append(text)
        else:
            rows = [[cell.text.replace("\n", "<br>").strip() for cell in row.cells] for row in item.rows]
            if not rows:
                continue
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + "---|" * len(rows[0]))
            out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
        out.append("")
    return "\n".join(out).strip() + "\n"


def read_source(path: str | Path) -> str:
    """Đọc tài liệu nguồn thành văn bản/Markdown cho LLM."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return docx_to_markdown(path)
    if suffix in {".md", ".txt", ".csv", ".json", ".yaml", ".yml"}:
        return path.read_text(encoding="utf-8")
    try:
        from markitdown import MarkItDown  # type: ignore[import-not-found]
    except ImportError as err:
        raise RuntimeError(
            f"Cần MarkItDown để đọc {suffix}: pip install 'wordreport[markitdown]' "
            "hoặc dùng MCP server markitdown (convert_to_markdown)."
        ) from err
    return MarkItDown().convert(str(path)).text_content


def outline(path: str | Path) -> list[dict]:
    document = Document(str(path))
    items = []
    for item in iter_block_items(document):
        if isinstance(item, Paragraph):
            level = _heading_level(item)
            if level and item.text.strip():
                items.append({"level": level, "text": item.text.strip()})
    return items


@dataclass
class LintIssue:
    severity: str  # error | warning | info
    rule: str
    message: str
    location: str = ""


_NUM_PREFIX = re.compile(r"^((?:[IVXLC]+\.)|(?:\d+(?:\.\d+)*\.?))\s")


def lint_docx(path: str | Path, profile: str = "hcmus-clc") -> list[LintIssue]:
    prof = load_profile(profile)
    allowed_fonts = {prof["fonts"]["body"], prof["fonts"]["code"], None}
    labels = (prof["caption"]["figure_label"], prof["caption"]["table_label"])
    document = Document(str(path))
    issues: list[LintIssue] = []
    items = list(iter_block_items(document))

    body_xml = document.element.body.xml
    if "TOC \\o" not in body_xml and " TOC " not in body_xml:
        issues.append(LintIssue("warning", "no-toc", "Tài liệu chưa có mục lục tự động (field TOC)."))
    has_page = any(" PAGE " in s.footer._element.xml or "PAGE" in s.footer._element.xml for s in document.sections)
    if not has_page and prof.get("footer", {}).get("page_number", "center") != "none":
        issues.append(LintIssue("warning", "no-page-number", "Footer chưa có số trang tự động (field PAGE)."))

    for label in find_watermarks(path):
        issues.append(LintIssue("warning", "ai-label", f"Còn nhãn trình tạo/AI: {label} - chạy watermark-remover."))

    from .barem import barem_outline_labels  # import muộn: barem.py dùng LintIssue của module này

    headings = [p.text.strip().casefold() for p in items
                if isinstance(p, Paragraph) and _heading_level(p) and p.text.strip()]
    for required in barem_outline_labels(prof):
        if not any(required.casefold() in h for h in headings):
            issues.append(LintIssue("warning", "barem-missing", f"Thiếu phần bắt buộc của barem: '{required}'."))

    fonts_seen: dict[str, int] = {}
    empty_run = 0
    prev_level = 0
    prev_numbers: dict[int, str] = {}
    # section đầu tiên của tài liệu nhiều section là trang bìa: logo/bảng thông tin không cần chú thích
    in_cover = len(document.sections) > 1
    for idx, item in enumerate(items):
        loc = f"khối #{idx}"
        cover_item = in_cover
        if isinstance(item, Paragraph) and item._p.pPr is not None and item._p.pPr.find(qn("w:sectPr")) is not None:
            in_cover = False
        if isinstance(item, Table):
            if cover_item:
                continue  # bảng bố cục/thông tin trang bìa
            borders = item._tbl.tblPr.find(qn("w:tblBorders"))
            if borders is not None and borders.find(qn("w:top")) is not None and borders.find(qn("w:top")).get(qn("w:val")) == "nil":
                continue  # bảng bố cục không viền (ví dụ khối thông tin trang bìa)
            prev = items[idx - 1] if idx > 0 else None
            if not (isinstance(prev, Paragraph) and _is_caption(prev, labels)):
                issues.append(LintIssue("info", "table-without-caption", "Bảng chưa có chú thích 'Bảng N: ...' phía trên.", loc))
            continue

        text = item.text
        level = _heading_level(item)
        for run in item.runs:
            if run.font.name not in allowed_fonts:
                fonts_seen[run.font.name] = fonts_seen.get(run.font.name, 0) + 1

        if not text.strip() and not _has_image(item):
            empty_run += 1
            if empty_run == 3:
                issues.append(LintIssue("info", "empty-paragraphs", "Có ≥3 đoạn trống liên tiếp - nên dùng spacing/page break thay vì Enter.", loc))
            continue
        empty_run = 0

        if level:
            style_name = item.style.name if item.style is not None else ""
            if style_name.startswith("Heading"):
                if level > prev_level + 1 and prev_level:
                    issues.append(LintIssue("warning", "heading-skip", f"Tiêu đề nhảy cấp H{prev_level} → H{level}: '{text[:60]}'", loc))
                match = _NUM_PREFIX.match(text.strip())
                if match:
                    prev_numbers[level] = match.group(1)
                elif level >= 2 and prev_numbers:
                    issues.append(LintIssue("info", "heading-unnumbered", f"Tiêu đề chưa đánh số trong khi các tiêu đề khác có số: '{text[:60]}'", loc))
                prev_level = level
        elif not _is_list_item(item):
            stripped = text.lstrip()
            if re.match(r"^[-•+\uf0b7\uf0a7]\s", stripped):
                issues.append(LintIssue("info", "manual-bullet", f"Gạch đầu dòng gõ tay, nên dùng danh sách tự động: '{stripped[:50]}'", loc))
            words = len(text.split())
            if words > 250:
                issues.append(LintIssue("info", "long-paragraph", f"Đoạn văn dài {words} từ - cân nhắc tách đoạn.", loc))

        if _has_image(item) and not cover_item:
            nxt = items[idx + 1] if idx + 1 < len(items) else None
            if not (isinstance(nxt, Paragraph) and _is_caption(nxt, labels)):
                issues.append(LintIssue("warning", "figure-without-caption", "Hình chưa có chú thích 'Hình N: ...' phía dưới.", loc))

        style_name = item.style.name if item.style is not None else ""
        if "  " in text.strip() and style_name != "Code Block":
            issues.append(LintIssue("info", "double-space", f"Có khoảng trắng kép: '{text.strip()[:50]}'", loc))
        for wrong, right in COMMON_TYPOS.items():
            if wrong in text:
                issues.append(LintIssue("warning", "typo", f"'{wrong}' → '{right}'", loc))

    for font, count in sorted(fonts_seen.items(), key=lambda kv: -kv[1]):
        issues.append(LintIssue("warning", "font-mismatch", f"Font '{font}' ({count} run) khác font chuẩn '{prof['fonts']['body']}'."))
    return issues


def lint_report(path: str | Path, profile: str = "hcmus-clc") -> dict:
    issues = lint_docx(path, profile)
    return {
        "file": str(path),
        "errors": sum(i.severity == "error" for i in issues),
        "warnings": sum(i.severity == "warning" for i in issues),
        "infos": sum(i.severity == "info" for i in issues),
        "issues": [asdict(i) for i in issues],
    }
