from docx import Document

from wordreport.inspector import lint_docx, outline
from wordreport.numbering import HeadingNumberer, to_roman
from wordreport.renderer import DocxRenderer
from wordreport.spec import ReportSpec
from wordreport.style_profile import _deep_merge, list_profiles, load_profile
from wordreport.watermark_remover import find_watermarks


def test_roman_and_heading_numbers():
    assert [to_roman(n) for n in (1, 4, 9, 14, 40)] == ["I", "IV", "IX", "XIV", "XL"]
    n = HeadingNumberer()
    assert [n.next(1), n.next(2), n.next(3), n.next(3), n.next(2), n.next(1), n.next(2)] == [
        "I.", "1.", "1.1.", "1.2.", "2.", "II.", "1."]


def test_profiles_inherit():
    assert {"hcmus-clc", "nd30-a4"} <= set(list_profiles())
    a4 = load_profile("nd30-a4")
    assert a4["page"]["size"] == "a4" and a4["page"]["margin_left_cm"] == 3.0
    assert a4["fonts"]["body"] == "Times New Roman"  # kế thừa từ hcmus-clc
    assert a4["headings"]["level1"]["bold"] is True


def test_render_example(tmp_path, example_spec):
    out = DocxRenderer(example_spec.profile, base_dir=tmp_path).save(example_spec, tmp_path / "r.docx")
    doc = Document(str(out))

    assert len(doc.sections) == 3
    header_text = doc.sections[2].header.paragraphs[0].text
    assert header_text == "Khoa Công nghệ Thông tin\t22CLC03"
    assert doc.sections[1].header.paragraphs[0].text == ""

    heads = [h["text"] for h in outline(out)]
    assert "I. Giới thiệu chung" in heads
    assert "2. Bài 2 – Web, DNS, DHCP Relay và RIP" in heads
    assert "2.2. Bài làm" in heads
    assert heads[-1] == "III. Tài liệu tham khảo"

    xml = doc.element.body.xml
    assert 'TOC \\o "1-3"' in xml
    assert xml.count("SEQ Hình") == 4 and xml.count("SEQ Bảng") == 3
    assert "PAGE" in doc.sections[2].footer._element.xml

    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption" and p.text.startswith(("Hình", "Bảng"))]
    assert "Bảng 1: Danh sách thành viên nhóm" in captions
    assert "Hình 4: Truy cập www.congtyhanhoa.vn từ PC3 và PC5" in captions

    issues = [i for i in lint_docx(out) if i.severity != "info"]
    assert issues == []

    # watermark-remover chạy mặc định: không còn nhãn trình tạo/AI, tác giả là thành viên nhóm
    assert find_watermarks(out) == []
    props = doc.core_properties
    assert props.author == "Nguyễn Minh An, Lê Thu Bình" and props.comments == ""
    assert props.created.year >= 2024


def test_inline_markup_and_missing_image(tmp_path):
    spec = ReportSpec.model_validate({
        "include_toc": False,
        "body": [
            {"type": "heading", "level": 1, "text": "A"},
            {"type": "paragraph", "text": "Chữ **đậm**, *nghiêng*, `code` và [link](https://example.com)."},
            {"type": "image", "path": "khong-ton-tai.png", "caption": "Ảnh thiếu"},
        ],
    })
    out = DocxRenderer("hcmus-clc", base_dir=tmp_path).save(spec, tmp_path / "i.docx")
    doc = Document(str(out))
    para = next(p for p in doc.paragraphs if p.text.startswith("Chữ"))
    runs = {r.text: r for r in para.runs}
    assert runs["đậm"].bold and runs["nghiêng"].italic
    assert runs["code"].font.name == "Consolas"
    assert "w:hyperlink" in para._p.xml
    assert any("CHÈN HÌNH" in p.text and "khong-ton-tai.png" in p.text for p in doc.paragraphs)


def test_lists_restart_numbering(tmp_path):
    spec = ReportSpec.model_validate({"include_toc": False, "body": [
        {"type": "list", "style": "number", "items": ["a", "b"]},
        {"type": "list", "style": "number", "items": ["c"]},
    ]})
    out = DocxRenderer().save(spec, tmp_path / "l.docx")
    doc = Document(str(out))
    num_ids = [p._p.pPr.numPr.numId.val for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert len(num_ids) == 3 and num_ids[0] == num_ids[1] != num_ids[2]


def _profile(**overrides):
    """Profile barem + ghi đè cho test các tuỳ chọn chung của renderer."""
    return _deep_merge(load_profile("hcmus-clc"), overrides)


def test_barem_skeleton_labels_and_level4(tmp_path):
    spec = ReportSpec.model_validate({
        "meta": {"subject": "Mạng máy tính", "instructors": ["GV"], "class_code": "X",
                 "members": [{"name": "An", "student_id": "1"}, {"name": "Bình", "student_id": "2"}]},
        "preface": ["Mở đầu."],
        "assignments": [{"member": "An", "tasks": ["Làm A", "Viết B"], "completion": "100%"}],
        "introduction": [{"type": "heading", "level": 1, "text": "Tóm tắt đề tài"},
                         {"type": "paragraph", "text": "Tóm tắt."}],
        "body": [
            {"type": "heading", "level": 1, "text": "Tổng quan", "label": "ch1"},
            {"type": "heading", "level": 2, "text": "Đặc tả"},
            {"type": "heading", "level": 3, "text": "Euler", "label": "euler"},
            {"type": "table", "columns": ["a"], "rows": [["1"]], "caption": "Kế hoạch", "label": "plan"},
            {"type": "figure_placeholder", "caption": "Sóng", "label": "wave"},
            {"type": "paragraph", "text": "Theo [[plan]] và [[wave]], xem mục [[euler]] của phần [[ch1]]."},
            {"type": "list", "style": "bullet", "items": ["ý"]},
        ],
        "references": [{"text": "Sách A"}],
    })
    out = DocxRenderer(spec.profile).save(spec, tmp_path / "b.docx")
    doc = Document(str(out))
    heads = [h["text"] for h in outline(out)]
    assert heads == ["LỜI MỞ ĐẦU", "MỤC LỤC", "I. Giới thiệu chung", "1. Thành viên nhóm",
                     "2. Bảng phân công công việc", "3. Tóm tắt đề tài", "II. Nội dung", "1. Tổng quan",
                     "1.1. Đặc tả", "1.1.1. Euler", "III. Tài liệu tham khảo"]
    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    assert captions[:3] == ["Bảng 1: Danh sách thành viên nhóm", "Bảng 2: Phân công công việc và mức độ hoàn thành",
                            "Bảng 3: Kế hoạch"]
    assert any(p.text == "Theo Bảng 3 và Hình 1, xem mục 1.1.1 của phần 1." for p in doc.paragraphs)
    members, tasks = doc.tables[1], doc.tables[2]  # bảng 0 là khối GVHD/nhóm trên bìa
    assert [c.text for c in members.rows[0].cells] == ["STT", "MSSV", "Họ và tên"]  # không có email -> bỏ cột
    assert tasks.rows[1].cells[2].text == "-\u00a0Làm A\n-\u00a0Viết B"
    toc = [p.text.split("\t")[0] for p in doc.paragraphs if p.style.name.startswith("toc")]
    assert "1.1.1. Euler" not in toc and "1.1. Đặc tả" in toc  # mục lục chỉ cấp 1-3
    numbering = doc.part.numbering_part.element.xml
    assert 'w:lvlText w:val="-"' in numbering  # barem: bullet -> gạch đầu dòng
    assert DocxRenderer("hcmus-clc").outline(spec)[2] == (1, "I. Giới thiệu chung")


def test_barem_check_reports_missing_parts():
    from wordreport.barem import check_spec

    spec = ReportSpec.model_validate({"body": [{"type": "paragraph", "text": "Xem [[khong-co]]."}]})
    messages = " | ".join(i.message for i in check_spec(spec))
    for expected in ("Thiếu LỜI MỞ ĐẦU", "chưa có thành viên", "Bảng phân công", "cần ít nhất một chương",
                     "Tài liệu tham khảo", "[[khong-co]]", "giảng viên hướng dẫn"):
        assert expected in messages


def test_chapter_numbering_front_matter_and_lists_of_figures(tmp_path):
    prof = _profile(barem={"enabled": False}, heading_numbering="chapter", references_numbered=False,
                    page={"border": {"style": "thickThinSmallGap", "color": "1F3864", "size": 24, "all_pages": True}},
                    footer={"page_number": "none"}, header={"left": "", "right": ""},
                    labels={"preface": "LỜI NÓI ĐẦU", "references": "TÀI LIỆU THAM KHẢO"},
                    caption={"separator": ". ", "table": {"bold_text": True, "italic": False}})
    spec = ReportSpec.model_validate({
        "preface": ["Lời nói đầu."],
        "front_matter": [{"type": "heading", "level": 1, "text": "TÓM TẮT", "numbered": False},
                         {"type": "paragraph", "text": "Tóm tắt."}],
        "list_of_figures": True, "list_of_tables": True,
        "body": [
            {"type": "heading", "level": 1, "text": "MỞ ĐẦU", "numbered": False},
            {"type": "heading", "level": 1, "text": "Tổng quan"},
            {"type": "heading", "level": 2, "text": "Giới thiệu"},
            {"type": "table", "columns": ["a"], "rows": [["1"]], "caption": "Bảng một"},
            {"type": "heading", "level": 1, "text": "Thiết kế"},
            {"type": "heading", "level": 2, "text": "Mức transistor"},
            {"type": "heading", "level": 3, "text": "Euler"},
            {"type": "figure_placeholder", "caption": "Hình một"},
            {"type": "figure_placeholder", "caption": "Hình hai"},
        ],
        "references": [{"text": "Sách A"}],
    })
    out = DocxRenderer(prof).save(spec, tmp_path / "c.docx")
    doc = Document(str(out))
    heads = [h["text"] for h in outline(out)]
    assert heads[:4] == ["LỜI NÓI ĐẦU", "TÓM TẮT", "MỤC LỤC", "DANH MỤC HÌNH"]
    assert {"MỞ ĐẦU", "I. Tổng quan", "1.1. Giới thiệu", "II. Thiết kế", "2.1. Mức transistor",
            "2.1.1. Euler", "TÀI LIỆU THAM KHẢO"} <= set(heads)
    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    assert captions == ["Bảng 1.1. Bảng một", "Hình 2.1. Hình một", "Hình 2.2. Hình hai"]
    lists = [p.text.rstrip("\t") for p in doc.paragraphs if p.style.name == "table of figures"]
    assert lists == ["Hình 2.1. Hình một", "Hình 2.2. Hình hai", "Bảng 1.1. Bảng một"]
    assert 'TOC \\h \\z \\c "Hình"' in doc.element.body.xml
    assert all("thickThinSmallGap" in sec._sectPr.xml for sec in doc.sections)
    assert "PAGE" not in doc.sections[2].footer._element.xml
    assert not any(p.text == "" and 'w:type="page"' in p._p.xml for p in doc.paragraphs)


def test_level1_format_global_captions_table_caption_below(tmp_path):
    prof = _profile(barem={"enabled": False}, heading_numbering="chapter", caption_numbering="global",
                    headings={"level1_format": "CHƯƠNG {n}:"}, caption={"table_position": "below"},
                    page={"number_from_front": True}, toc={"include_self": False})
    spec = ReportSpec.model_validate({
        "meta": {"topic": "Dòng một\nDòng hai"},
        "body": [
            {"type": "heading", "level": 1, "text": "GIỚI THIỆU"},
            {"type": "heading", "level": 2, "text": "Khái niệm"},
            {"type": "paragraph", "text": "V~DROP~ = x^2^"},
            {"type": "figure_placeholder", "caption": "Hình A"},
            {"type": "heading", "level": 1, "text": "THIẾT KẾ"},
            {"type": "table", "columns": ["a", "b"], "rows": [["1", "2"]], "caption": "Bảng A"},
            {"type": "figure_placeholder", "caption": "Hình B"},
        ],
    })
    out = DocxRenderer(prof).save(spec, tmp_path / "l.docx")
    doc = Document(str(out))
    heads = [h["text"] for h in outline(out)]
    assert "CHƯƠNG 1: GIỚI THIỆU" in heads and "1.1. Khái niệm" in heads and "CHƯƠNG 2: THIẾT KẾ" in heads
    toc = [p.text for p in doc.paragraphs if p.style.name.startswith("toc")]
    assert not any(t.startswith("MỤC LỤC") for t in toc)
    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    assert captions == ["Hình 1: Hình A", "Bảng 1: Bảng A", "Hình 2: Hình B"]
    body = list(doc.element.body)
    caption_p = next(p for p in doc.paragraphs if p.text == "Bảng 1: Bảng A")
    assert body.index(caption_p._p) == body.index(doc.tables[-1]._tbl) + 1
    eq = next(p for p in doc.paragraphs if p.text == "VDROP = x2")
    runs = {r.text: r for r in eq.runs}
    assert runs["DROP"].font.subscript and runs["2"].font.superscript
    assert 'w:start="1"' in doc.sections[1]._sectPr.xml and "w:start=" not in doc.sections[2]._sectPr.xml
    assert doc.core_properties.subject == "Dòng một Dòng hai"
