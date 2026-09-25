from docx import Document

from wordreport.inspector import docx_to_markdown, lint_docx


def _bad_doc(path):
    doc = Document()
    doc.add_heading("I. Giới thiệu", level=1)
    doc.add_heading("Chi tiết", level=3)
    doc.add_paragraph("- gạch đầu dòng gõ tay")
    doc.add_paragraph("TP. Hồ Chí Mình – 2023")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("")
    run = doc.add_paragraph().add_run("Arial text")
    run.font.name = "Arial"
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "x"
    doc.save(str(path))


def test_lint_detects_common_problems(tmp_path):
    path = tmp_path / "bad.docx"
    _bad_doc(path)
    rules = {i.rule for i in lint_docx(path)}
    assert {"no-toc", "no-page-number", "heading-skip", "manual-bullet", "typo",
            "empty-paragraphs", "font-mismatch", "table-without-caption"} <= rules


def test_docx_to_markdown(tmp_path):
    path = tmp_path / "bad.docx"
    _bad_doc(path)
    md = docx_to_markdown(path)
    assert "# I. Giới thiệu" in md and "### Chi tiết" in md and "| x |" in md
